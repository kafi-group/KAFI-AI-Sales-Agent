"""Module for comparing AI-enriched spreadsheet files against Sales Agent DB contacts.

Provides read-only gap reports and safe-fill merges (only populating missing/blank fields,
preserving existing non-empty DB values 100%).
"""

from __future__ import annotations

import csv
import io
from typing import Any
import openpyxl
from sqlalchemy.orm import Session, joinedload
from db.models import Buyer, Contact, AppUser
from modules.buyers import normalize_buyer_key as _normalize_name, buyer_website_domain as _dedupe_domain

COMPARE_COLUMNS = [
    ("company_name", "Company Name"),
    ("contact_name", "Contact Person / Name"),
    ("designation", "Designation / Title"),
    ("email", "Email Address #1"),
    ("secondary_email", "Email Address #2"),
    ("phone", "Phone Number #1"),
    ("secondary_phone", "Phone Number #2 / Mobile"),
    ("website_url", "Website URL"),
    ("country", "Country"),
    ("city", "City"),
    ("address", "Address"),
    ("linkedin_url", "LinkedIn URL"),
    ("facebook_url", "Facebook URL"),
    ("instagram_url", "Instagram URL"),
    ("remarks", "Remarks / Notes"),
]


def _parse_uploaded_file(file_content: bytes, filename: str) -> list[dict[str, Any]]:
    """Parse uploaded bytes into a list of row dicts."""
    fname = filename.lower()
    raw_records: list[dict[str, Any]] = []

    if fname.endswith(".csv"):
        text = file_content.decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(text))
        for r in reader:
            raw_records.append({str(k).strip(): str(v).strip() for k, v in r.items() if k is not None})
    elif fname.endswith((".xlsx", ".xlsm")):
        wb = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True, read_only=True)
        max_rows = -1
        best_records: list[dict[str, Any]] = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            sheet_rows = list(ws.iter_rows(values_only=True))
            if not sheet_rows:
                continue
            header_row = [str(c).strip() if c is not None else "" for c in sheet_rows[0]]
            cols_lower = [c.lower() for c in header_row]
            has_contacts = any("company" in c or "buyer" in c or "contact" in c for c in cols_lower)
            records = []
            for r_vals in sheet_rows[1:]:
                if not any(r_vals):
                    continue
                r_dict = {}
                for i, h in enumerate(header_row):
                    if h:
                        v = r_vals[i] if i < len(r_vals) else ""
                        r_dict[h] = str(v).strip() if v is not None else ""
                records.append(r_dict)

            if (has_contacts and len(records) > max_rows) or (not best_records and records):
                max_rows = len(records)
                best_records = records

        raw_records = best_records
    elif fname.endswith(".xls"):
        import xlrd
        book = xlrd.open_workbook(file_contents=file_content)
        sheet = book.sheet_by_index(0)
        header_row = [str(sheet.cell_value(0, c)).strip() for c in range(sheet.ncols)]
        for r in range(1, sheet.nrows):
            r_dict = {}
            for c in range(sheet.ncols):
                h = header_row[c]
                if h:
                    v = sheet.cell_value(r, c)
                    r_dict[h] = str(v).strip() if v is not None else ""
            raw_records.append(r_dict)
    else:
        raise ValueError("Unsupported file format. Please upload an Excel (.xlsx) or CSV file.")

    rows = raw_records
    normalized_rows = []
    for r in rows:
        norm = {}
        for k, v in r.items():
            k_clean = str(k).strip().lower().replace(" ", "_").replace("#", "")
            val_clean = str(v).strip() if v is not None else ""
            if val_clean in ("nan", "none", "null"):
                val_clean = ""
            norm[k_clean] = val_clean
            norm[str(k).strip()] = val_clean

        company = (
            norm.get("company_name")
            or norm.get("company")
            or norm.get("buyer_name")
            or norm.get("name")
            or ""
        )
        contact = (
            norm.get("contact_name")
            or norm.get("contact_person")
            or norm.get("person_name")
            or norm.get("contact")
            or ""
        )
        designation = (
            norm.get("designation")
            or norm.get("title")
            or norm.get("job_title")
            or norm.get("role")
            or ""
        )
        email1 = (
            norm.get("email")
            or norm.get("email_address")
            or norm.get("email_1")
            or norm.get("primary_email")
            or norm.get("primary_email_no.")
            or ""
        )
        email2 = (
            norm.get("secondary_email")
            or norm.get("email_2")
            or norm.get("alternate_email")
            or norm.get("secondary_email_no.")
            or ""
        )
        phone1 = (
            norm.get("phone")
            or norm.get("phone_number")
            or norm.get("mobile")
            or norm.get("primary_mobile_no.")
            or norm.get("primary_phone_no.")
            or norm.get("contact_1")
            or ""
        )
        phone2 = (
            norm.get("secondary_phone")
            or norm.get("secondary_mobile")
            or norm.get("secondary_mobile_no.")
            or norm.get("secondary_phone_no.")
            or norm.get("contact_2")
            or norm.get("phone_2")
            or ""
        )
        website = (
            norm.get("website_url")
            or norm.get("website")
            or norm.get("domain")
            or norm.get("url")
            or ""
        )
        country = norm.get("country") or norm.get("location") or ""
        city = norm.get("city") or ""
        address = norm.get("address") or norm.get("street") or ""
        linkedin = norm.get("linkedin_url") or norm.get("linkedin") or norm.get("linkedin_company_url") or ""
        facebook = norm.get("facebook_url") or norm.get("facebook") or norm.get("facebook_company_url") or ""
        instagram = norm.get("instagram_url") or norm.get("instagram") or norm.get("instagram_company_url") or ""
        remarks = norm.get("remarks") or norm.get("notes") or norm.get("comment") or ""
        buyer_id = norm.get("buyer_id") or norm.get("id") or ""
        company_grading = (
            norm.get("company_grading")
            or norm.get("excel_grading")
            or norm.get("companies_grading")
            or norm.get("grading")
            or norm.get("company_grading_/_excel_grading")
            or norm.get("excel_/_company_grading")
            or ""
        )
        product_interest = norm.get("product_interest") or norm.get("product") or norm.get("products") or ""
        industry = norm.get("industry") or norm.get("business_type") or ""

        row_dict = {
            "buyer_id": buyer_id,
            "company_name": company,
            "contact_name": contact,
            "designation": designation,
            "email": email1,
            "secondary_email": email2,
            "phone": phone1,
            "secondary_phone": phone2,
            "website_url": website,
            "country": country,
            "city": city,
            "address": address,
            "company_grading": company_grading,
            "product_interest": product_interest,
            "industry": industry,
            "linkedin_url": linkedin,
            "facebook_url": facebook,
            "instagram_url": instagram,
            "remarks": remarks,
            "_raw": r,
        }
        normalized_rows.append(row_dict)
    return normalized_rows


def _get_db_field_val(buyer: Buyer, field_key: str) -> str:
    """Safely extract field value from Buyer or primary Contact."""
    if field_key == "company_name":
        return buyer.company_name or ""
    elif field_key == "website_url":
        return buyer.website_url or ""
    elif field_key == "country":
        return buyer.country or ""
    elif field_key == "city":
        return buyer.city or ""
    elif field_key == "address":
        return buyer.address or ""
    elif field_key == "remarks":
        return buyer.remarks or ""
    elif field_key == "facebook_url":
        return buyer.facebook_company_url or ""
    elif field_key == "instagram_url":
        return buyer.instagram_company_url or ""
    elif field_key == "linkedin_url":
        val = buyer.linkedin_company_url or ""
        if not val and buyer.contacts:
            val = buyer.contacts[0].linkedin_profile_url or ""
        return val

    # Contact-level fields
    primary_contact = buyer.contacts[0] if buyer.contacts else None
    if not primary_contact:
        return ""

    if field_key == "contact_name":
        return primary_contact.full_name or ""
    elif field_key == "designation":
        return primary_contact.designation or ""
    elif field_key == "email":
        return primary_contact.email or ""
    elif field_key == "secondary_email":
        return primary_contact.secondary_email or ""
    elif field_key == "phone":
        return primary_contact.phone or primary_contact.primary_phone or ""
    elif field_key == "secondary_phone":
        return primary_contact.secondary_phone or primary_contact.secondary_mobile or ""

    return ""


def generate_enrichment_comparison_report(
    db: Session,
    *,
    file_content: bytes,
    filename: str,
    user_id: int | None = None,
    table_source: str = "master_table",
    master_type: str = "fmcg",
) -> dict[str, Any]:
    """Read-only analysis comparing uploaded enriched file against Sales Agent contacts."""
    uploaded_rows = _parse_uploaded_file(file_content, filename)

    query = db.query(Buyer).options(joinedload(Buyer.contacts))
    if table_source == "old_clients":
        query = query.filter(Buyer.source == "old_clients")
    elif table_source == "master_table":
        query = query.filter(Buyer.source != "old_clients")
        if master_type:
            query = query.filter(Buyer.master_type == master_type)

    if user_id is not None:
        query = query.filter(Buyer.assigned_to_user_id == user_id)
        user_obj = db.query(AppUser).filter(AppUser.id == user_id).first()
        user_name = user_obj.full_name if user_obj else f"User #{user_id}"
    else:
        user_name = "All Assigned Contacts"

    db_buyers = query.all()
    db_total_contacts = len(db_buyers)

    by_name = {}
    by_domain = {}
    by_email = {}
    by_phone = {}

    for buyer in db_buyers:
        if buyer.company_name:
            by_name[_normalize_name(buyer.company_name)] = buyer
        if buyer.website_url:
            dom = _dedupe_domain(buyer.website_url)
            if dom:
                by_domain[dom] = buyer

        for c in buyer.contacts or []:
            if c.email:
                by_email[c.email.strip().lower()] = buyer
            if c.secondary_email:
                by_email[c.secondary_email.strip().lower()] = buyer
            if c.phone:
                by_phone[c.phone.strip()] = buyer
            if c.primary_phone:
                by_phone[c.primary_phone.strip()] = buyer

    matched_pairs: list[tuple[dict[str, Any], Buyer]] = []
    unmatched_rows: list[dict[str, Any]] = []

    for row in uploaded_rows:
        comp_name = row.get("company_name") or ""
        domain = _dedupe_domain(row.get("website_url"))
        email = (row.get("email") or "").strip().lower()
        phone = (row.get("phone") or "").strip()

        matched_buyer = None
        if comp_name and _normalize_name(comp_name) in by_name:
            matched_buyer = by_name[_normalize_name(comp_name)]
        elif domain and domain in by_domain:
            matched_buyer = by_domain[domain]
        elif email and email in by_email:
            matched_buyer = by_email[email]
        elif phone and phone in by_phone:
            matched_buyer = by_phone[phone]

        if matched_buyer:
            matched_pairs.append((row, matched_buyer))
        else:
            unmatched_rows.append(row)

    column_analysis = []
    total_potential_new_fills = 0
    sample_fills = []

    for field_key, field_label in COMPARE_COLUMNS:
        db_populated = 0
        db_missing = 0
        file_populated = 0
        new_fill_potential = 0
        protected_existing = 0

        for file_row, buyer in matched_pairs:
            db_val = _get_db_field_val(buyer, field_key)
            db_has_val = bool(db_val and str(db_val).strip())
            file_val = file_row.get(field_key)
            file_has_val = bool(file_val and str(file_val).strip())

            if db_has_val:
                db_populated += 1
            else:
                db_missing += 1

            if file_has_val:
                file_populated += 1

            if not db_has_val and file_has_val:
                new_fill_potential += 1
                total_potential_new_fills += 1
                if len(sample_fills) < 6:
                    sample_fills.append(
                        {
                            "company_name": buyer.company_name,
                            "field_name": field_label,
                            "new_value": file_val,
                        }
                    )

            if (
                db_has_val
                and file_has_val
                and str(db_val).strip().lower() != str(file_val).strip().lower()
            ):
                protected_existing += 1

        column_analysis.append(
            {
                "field_key": field_key,
                "field_label": field_label,
                "db_populated_count": db_populated,
                "db_missing_count": db_missing,
                "file_populated_count": file_populated,
                "new_fill_count": new_fill_potential,
                "protected_count": protected_existing,
            }
        )

    return {
        "user_name": user_name,
        "table_source": table_source,
        "filename": filename,
        "db_total_contacts": db_total_contacts,
        "uploaded_file_contacts": len(uploaded_rows),
        "matched_contacts_count": len(matched_pairs),
        "unmatched_contacts_count": len(unmatched_rows),
        "total_potential_new_fills": total_potential_new_fills,
        "column_analysis": column_analysis,
        "sample_fills": sample_fills,
    }


def execute_safe_fill_merge(
    db: Session,
    *,
    file_content: bytes,
    filename: str,
    user_id: int | None = None,
    table_source: str = "master_table",
    master_type: str = "fmcg",
) -> dict[str, Any]:
    """Perform a Safe Merge: Populates missing/blank fields only, 100% preserving existing DB values."""
    uploaded_rows = _parse_uploaded_file(file_content, filename)

    query = db.query(Buyer).options(joinedload(Buyer.contacts))
    if table_source == "old_clients":
        query = query.filter(Buyer.source == "old_clients")
    elif table_source == "master_table":
        query = query.filter(Buyer.source != "old_clients")
        if master_type:
            query = query.filter(Buyer.master_type == master_type)

    if user_id is not None:
        query = query.filter(Buyer.assigned_to_user_id == user_id)

    db_buyers = query.all()

    by_id = {buyer.id: buyer for buyer in db_buyers}
    by_name = {}
    by_domain = {}
    by_email = {}
    by_phone = {}

    for buyer in db_buyers:
        if buyer.company_name:
            by_name[_normalize_name(buyer.company_name)] = buyer
        if buyer.website_url:
            dom = _dedupe_domain(buyer.website_url)
            if dom:
                by_domain[dom] = buyer
        for c in buyer.contacts or []:
            if c.email:
                by_email[c.email.strip().lower()] = buyer
            if c.secondary_email:
                by_email[c.secondary_email.strip().lower()] = buyer
            if c.phone:
                by_phone[c.phone.strip()] = buyer
            if c.primary_phone:
                by_phone[c.primary_phone.strip()] = buyer

    contacts_updated = 0
    total_fields_filled = 0
    protected_fields_count = 0

    buyer_field_map = {
        "country": "country",
        "city": "city",
        "address": "address",
        "website_url": "website_url",
        "company_grading": "company_grading",
        "product_interest": "product_interest",
        "industry": "industry",
        "remarks": "remarks",
        "facebook_url": "facebook_company_url",
        "instagram_url": "instagram_company_url",
        "linkedin_url": "linkedin_company_url",
    }

    contact_field_map = {
        "contact_name": "full_name",
        "designation": "designation",
        "email": "email",
        "secondary_email": "secondary_email",
        "phone": "phone",
        "secondary_phone": "secondary_phone",
        "linkedin_url": "linkedin_profile_url",
    }

MAX_FIELD_LENGTHS = {
    "country": 100,
    "city": 255,
    "company_grading": 50,
    "website_url": 512,
    "industry": 255,
    "product_interest": 512,
    "full_name": 255,
    "designation": 255,
    "email": 255,
    "secondary_email": 255,
    "phone": 50,
    "secondary_phone": 50,
    "primary_phone": 50,
    "linkedin_profile_url": 512,
    "linkedin_company_url": 512,
    "facebook_company_url": 512,
    "instagram_company_url": 512,
}

    for row in uploaded_rows:
        buyer_id_raw = str(row.get("buyer_id") or "").strip()
        comp_name = row.get("company_name") or ""
        domain = _dedupe_domain(row.get("website_url"))
        email = (row.get("email") or "").strip().lower()
        phone = (row.get("phone") or "").strip()

        buyer = None
        if buyer_id_raw.isdigit() and int(buyer_id_raw) in by_id:
            buyer = by_id[int(buyer_id_raw)]
        elif comp_name and _normalize_name(comp_name) in by_name:
            buyer = by_name[_normalize_name(comp_name)]
        elif domain and domain in by_domain:
            buyer = by_domain[domain]
        elif email and email in by_email:
            buyer = by_email[email]
        elif phone and phone in by_phone:
            buyer = by_phone[phone]

        if not buyer:
            continue

        buyer_modified = False

        # 1. Update Buyer model fields
        for file_key, db_attr in buyer_field_map.items():
            current_val = getattr(buyer, db_attr, None)
            new_val = row.get(file_key)
            if not current_val or not str(current_val).strip():
                if new_val and str(new_val).strip():
                    val_str = str(new_val).strip()
                    max_len = MAX_FIELD_LENGTHS.get(db_attr)
                    if max_len and len(val_str) > max_len:
                        val_str = val_str[:max_len].strip()
                    setattr(buyer, db_attr, val_str)
                    buyer_modified = True
                    total_fields_filled += 1
            else:
                if (
                    new_val
                    and str(new_val).strip()
                    and str(current_val).strip().lower() != str(new_val).strip().lower()
                ):
                    protected_fields_count += 1

        # 2. Update Primary Contact fields
        if not buyer.contacts:
            contact_name = row.get("contact_name") or buyer.company_name or "General Contact"
            primary_contact = Contact(buyer_id=buyer.id, full_name=str(contact_name)[:255])
            db.add(primary_contact)
            buyer.contacts.append(primary_contact)
            buyer_modified = True
        else:
            primary_contact = buyer.contacts[0]

        for file_key, db_attr in contact_field_map.items():
            current_val = getattr(primary_contact, db_attr, None)
            new_val = row.get(file_key)
            if not current_val or not str(current_val).strip():
                if new_val and str(new_val).strip():
                    val_str = str(new_val).strip()
                    max_len = MAX_FIELD_LENGTHS.get(db_attr)
                    if max_len and len(val_str) > max_len:
                        val_str = val_str[:max_len].strip()
                    setattr(primary_contact, db_attr, val_str)
                    if db_attr == "phone" and (not getattr(primary_contact, "primary_phone", None) or not str(primary_contact.primary_phone).strip()):
                        primary_contact.primary_phone = val_str[:50]
                    buyer_modified = True
                    total_fields_filled += 1
            else:
                if (
                    new_val
                    and str(new_val).strip()
                    and str(current_val).strip().lower() != str(new_val).strip().lower()
                ):
                    protected_fields_count += 1

        if buyer_modified:
            contacts_updated += 1

    db.commit()

    return {
        "status": "success",
        "contacts_updated": contacts_updated,
        "total_fields_filled": total_fields_filled,
        "protected_fields_count": protected_fields_count,
        "message": f"Successfully enriched {contacts_updated} contacts with {total_fields_filled} new missing fields filled! {protected_fields_count} existing fields were 100% protected.",
    }


AVAILABLE_REPORT_COLUMNS = [
    ("contact_name", "Contact Person / Name"),
    ("designation", "Designation / Title"),
    ("phone", "Primary Phone Number"),
    ("secondary_phone", "Secondary Phone / Mobile"),
    ("email", "Primary Email Address"),
    ("secondary_email", "Secondary Email Address"),
    ("website_url", "Website URL"),
    ("country", "Country"),
    ("city", "City"),
    ("address", "Address"),
    ("company_grading", "Company Grading"),
    ("product_interest", "Product / Remarks"),
]


def generate_missing_data_report(
    db: Session,
    *,
    section: str = "master",
    column_key: str = "contact_name",
    user_id: int | None = None,
    master_type: str = "fmcg",
    viewer: AppUser,
) -> dict[str, Any]:
    """Generates column-wise missing data report for specified section, column, and user scope."""
    # Scope resolution
    is_admin = getattr(viewer, "role", "") == "admin" or str(getattr(viewer, "role", "")).endswith("admin")
    if not is_admin:
        target_user_id = viewer.id
        user_name = viewer.full_name or viewer.username
    else:
        target_user_id = user_id
        if target_user_id:
            u_obj = db.query(AppUser).filter(AppUser.id == target_user_id).first()
            user_name = u_obj.full_name or u_obj.username if u_obj else f"User #{target_user_id}"
        else:
            user_name = "All Assigned Users"

    # Query setup with joinedload
    query = db.query(Buyer).options(joinedload(Buyer.contacts))

    # Section filtering
    sec_clean = (section or "master").strip().lower()
    if sec_clean == "old_clients":
        query = query.filter(Buyer.source == "old_clients")
        section_label = "Old Clients"
    elif sec_clean in ("new_search_lead", "discover"):
        query = query.filter(Buyer.intake_method == "discover")
        section_label = "New Search Lead"
    elif sec_clean in ("khalid_focused", "focused"):
        query = query.filter(Buyer.interested_clients_list_at.isnot(None))
        section_label = "Khalid Focused Sales"
    elif sec_clean == "master":
        query = query.filter(Buyer.source != "old_clients")
        if master_type:
            query = query.filter(Buyer.master_type == master_type)
        if master_type == "minerals_ores":
            section_label = "Master Table (Minerals & Ores)"
        elif master_type == "other_items":
            section_label = "Master Table (Other Items)"
        else:
            section_label = "Master Table (FMCG)"
    else:
        section_label = "All Sections Combined"

    # User scoping
    if target_user_id is not None:
        query = query.filter(Buyer.assigned_to_user_id == target_user_id)

    db_buyers = query.all()
    total_contacts = len(db_buyers)

    column_label_dict = dict(AVAILABLE_REPORT_COLUMNS)
    col_label = column_label_dict.get(column_key, column_key.replace("_", " ").title())

    missing_rows = []
    populated_count = 0

    for buyer in db_buyers:
        val = _get_db_field_val(buyer, column_key)
        has_val = bool(val and str(val).strip() and str(val).strip().lower() not in ("select...", "nan", "none", "null", "-"))

        if has_val:
            populated_count += 1
        else:
            contact = buyer.contacts[0] if buyer.contacts else None
            missing_rows.append({
                "id": buyer.id,
                "legacy_serial_no": buyer.legacy_serial_no or buyer.id,
                "company_name": buyer.company_name or "Unnamed Company",
                "assigned_to": buyer.assigned_to or (buyer.assigned_to_user_id and f"User #{buyer.assigned_to_user_id}") or "unassigned",
                "company_grading": buyer.company_grading or "-",
                "country": buyer.country or "-",
                "city": buyer.city or "-",
                "missing_column_key": column_key,
                "missing_column_label": col_label,
                "website_url": buyer.website_url or "-",
                "contact_person": contact.full_name if contact else "-",
                "primary_phone": (contact.phone or contact.primary_phone) if contact else "-",
                "primary_email": contact.email if contact else "-",
            })

    missing_count = len(missing_rows)
    missing_pct = round((missing_count / total_contacts * 100), 2) if total_contacts > 0 else 0.0

    return {
        "section": sec_clean,
        "section_label": section_label,
        "column_key": column_key,
        "column_label": col_label,
        "user_id": target_user_id,
        "user_name": user_name,
        "total_contacts": total_contacts,
        "missing_count": missing_count,
        "populated_count": populated_count,
        "missing_percentage": missing_pct,
        "available_columns": [{"key": k, "label": v} for k, v in AVAILABLE_REPORT_COLUMNS],
        "missing_rows": missing_rows,
    }

