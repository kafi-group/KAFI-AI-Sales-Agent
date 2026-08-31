"""Horeka B2B Price List module for line-item pricing, export, and client dispatch."""

from __future__ import annotations

import io
import json
import logging
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Optional

from sqlalchemy import desc, func, or_, select
from sqlalchemy.orm import Session

from db.models import HorekaLineItem
from modules.email_attachments import register_attachment_from_bytes, public_attachment

logger = logging.getLogger(__name__)

_BACKEND_DIR = Path(__file__).resolve().parent.parent
CATALOG_JSON_PATH = _BACKEND_DIR / "data" / "kafi_essence_catalog.json"
PRICE_TIERS_JSON_PATH = _BACKEND_DIR / "data" / "category_price_tiers.json"


def seed_horeka_line_items_if_needed(db: Session) -> int:
    """Pre-seed Horeka line items from kafi_essence_catalog.json if table is empty."""
    count = db.scalar(select(func.count(HorekaLineItem.id))) or 0
    if count > 0:
        return count

    if not CATALOG_JSON_PATH.exists():
        logger.warning("Catalog JSON not found at %s", CATALOG_JSON_PATH)
        return 0

    try:
        catalog_data = json.loads(CATALOG_JSON_PATH.read_text(encoding="utf-8"))
        pricing_data = {}
        if PRICE_TIERS_JSON_PATH.exists():
            pricing_data = json.loads(PRICE_TIERS_JSON_PATH.read_text(encoding="utf-8")).get("categories", {})
    except Exception as exc:
        logger.error("Error reading initial catalog data: %s", exc)
        return 0

    products = catalog_data.get("products", [])
    items_to_add: list[HorekaLineItem] = []

    for p in products:
        cat_key = (p.get("category") or "other").lower()
        pricing_info = pricing_data.get(cat_key, {})
        unit = pricing_info.get("unit", "USD/carton")
        std_price = float(pricing_info.get("standard", 18.0))
        bulk_tier1 = float(pricing_info.get("bulk_carton") or pricing_info.get("bulk_100mt") or (std_price * 0.92))
        bulk_tier2 = float(pricing_info.get("bulk_container") or pricing_info.get("bulk_500mt") or (std_price * 0.85))

        item = HorekaLineItem(
            sno=p.get("sno") or p.get("id") or 0,
            category=cat_key.replace("_", " ").title(),
            sub_category=None,
            brand=p.get("brand") or "ESSENCE",
            item_code=f"KAF-{cat_key[:3].upper()}-{p.get('sno') or p.get('id'):03d}",
            product_name=p.get("name") or "Product",
            packaging=p.get("packaging") or "Standard Export Carton",
            unit=unit,
            standard_price=round(std_price, 2),
            bulk_tier1_price=round(bulk_tier1, 2),
            bulk_tier2_price=round(bulk_tier2, 2),
            moq="50 Master Cartons" if "carton" in unit.lower() else "1 FCL / 25 MT",
            stock_status="In Stock",
            notes="Factory direct export pricing",
        )
        items_to_add.append(item)

    if items_to_add:
        db.add_all(items_to_add)
        db.commit()
        logger.info("Successfully seeded %d Horeka line items into database", len(items_to_add))
        return len(items_to_add)

    return 0


def list_horeka_items(
    db: Session,
    category: Optional[str] = None,
    search: Optional[str] = None,
    stock_status: Optional[str] = None,
    limit: int = 500,
    offset: int = 0,
) -> dict[str, Any]:
    """Retrieve Horeka line items with optional filtering and summary stats."""
    seed_horeka_line_items_if_needed(db)

    query = select(HorekaLineItem)
    if category and category.lower() != "all":
        query = query.where(HorekaLineItem.category.ilike(f"%{category}%"))
    if stock_status and stock_status.lower() != "all":
        query = query.where(HorekaLineItem.stock_status.ilike(f"%{stock_status}%"))
    if search and search.strip():
        term = f"%{search.strip()}%"
        query = query.where(
            or_(
                HorekaLineItem.product_name.ilike(term),
                HorekaLineItem.item_code.ilike(term),
                HorekaLineItem.packaging.ilike(term),
                HorekaLineItem.category.ilike(term),
            )
        )

    total = db.scalar(select(func.count()).select_from(query.subquery())) or 0
    items = db.scalars(query.order_by(HorekaLineItem.category, HorekaLineItem.sno, HorekaLineItem.id).offset(offset).limit(limit)).all()

    # Get distinct categories
    cats = db.scalars(select(HorekaLineItem.category).distinct().order_by(HorekaLineItem.category)).all()

    return {
        "total": total,
        "categories": cats,
        "items": [
            {
                "id": it.id,
                "sno": it.sno,
                "category": it.category,
                "sub_category": it.sub_category,
                "brand": it.brand,
                "item_code": it.item_code,
                "product_name": it.product_name,
                "packaging": it.packaging,
                "unit": it.unit,
                "standard_price": it.standard_price,
                "bulk_tier1_price": it.bulk_tier1_price,
                "bulk_tier2_price": it.bulk_tier2_price,
                "moq": it.moq,
                "stock_status": it.stock_status,
                "notes": it.notes,
                "updated_at": it.updated_at.isoformat() if it.updated_at else None,
            }
            for it in items
        ],
    }


def update_horeka_item(db: Session, item_id: int, payload: dict[str, Any]) -> dict[str, Any]:
    item = db.get(HorekaLineItem, item_id)
    if not item:
        raise ValueError(f"Horeka item #{item_id} not found")

    allowed_fields = {
        "category",
        "sub_category",
        "brand",
        "item_code",
        "product_name",
        "packaging",
        "unit",
        "standard_price",
        "bulk_tier1_price",
        "bulk_tier2_price",
        "moq",
        "stock_status",
        "notes",
    }
    for k, v in payload.items():
        if k in allowed_fields:
            setattr(item, k, v)

    item.updated_at = datetime.now(timezone.utc)
    db.commit()
    db.refresh(item)
    return {
        "id": item.id,
        "product_name": item.product_name,
        "standard_price": item.standard_price,
        "updated_at": item.updated_at.isoformat() if item.updated_at else None,
    }


def create_horeka_item(db: Session, payload: dict[str, Any]) -> dict[str, Any]:
    max_sno = db.scalar(select(func.max(HorekaLineItem.sno))) or 0
    item = HorekaLineItem(
        sno=payload.get("sno") or (max_sno + 1),
        category=payload.get("category", "General"),
        sub_category=payload.get("sub_category"),
        brand=payload.get("brand", "ESSENCE"),
        item_code=payload.get("item_code") or f"KAF-NEW-{max_sno + 1:03d}",
        product_name=payload.get("product_name", "New Item"),
        packaging=payload.get("packaging", "Standard Export Carton"),
        unit=payload.get("unit", "USD/carton"),
        standard_price=float(payload.get("standard_price", 0.0)),
        bulk_tier1_price=float(payload.get("bulk_tier1_price", 0.0)) if payload.get("bulk_tier1_price") is not None else None,
        bulk_tier2_price=float(payload.get("bulk_tier2_price", 0.0)) if payload.get("bulk_tier2_price") is not None else None,
        moq=payload.get("moq", "10 Master Cartons"),
        stock_status=payload.get("stock_status", "In Stock"),
        notes=payload.get("notes"),
    )
    db.add(item)
    db.commit()
    db.refresh(item)
    return {
        "id": item.id,
        "sno": item.sno,
        "product_name": item.product_name,
        "standard_price": item.standard_price,
    }


def delete_horeka_item(db: Session, item_id: int) -> bool:
    item = db.get(HorekaLineItem, item_id)
    if not item:
        return False
    db.delete(item)
    db.commit()
    return True


def generate_horeka_excel(db: Session) -> bytes:
    """Generate a formatted Excel sheet of the full Horeka B2B Price List."""
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter

        seed_horeka_line_items_if_needed(db)
        items = db.scalars(select(HorekaLineItem).order_by(HorekaLineItem.category, HorekaLineItem.sno, HorekaLineItem.id)).all()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Horeka B2B Price List"
        ws.views.sheetView[0].showGridLines = True

        # Header Title
        ws.merge_cells("A1:H1")
        ws["A1"] = "KAFI COMMODITIES (PVT.) LTD. — HOREKA B2B WHOLESALE PRICE LIST"
        ws["A1"].font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
        ws["A1"].fill = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid")
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 36

        # Subtitle
        ws.merge_cells("A2:H2")
        ws["A2"] = f"Brand: ESSENCE | Karachi Port (FOB/CNF) | Generated: {datetime.now().strftime('%d-%b-%Y')}"
        ws["A2"].font = Font(name="Arial", size=10, italic=True, color="1E293B")
        ws["A2"].fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 20

        # Table Column Headers
        headers = [
            "Item Code",
            "Category",
            "Product Description",
            "Packaging / Specifications",
            "Unit",
            "Standard B2B Price",
            "Bulk Tier 1 Price",
            "MOQ",
        ]

        ws.append([])  # Row 3 blank
        ws.append(headers)  # Row 4 headers
        ws.row_dimensions[4].height = 26

        header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=4, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        # Populate rows
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        for idx, it in enumerate(items, start=5):
            ws.append([
                it.item_code or f"KAF-{it.id:03d}",
                it.category,
                it.product_name,
                it.packaging,
                it.unit,
                it.standard_price,
                it.bulk_tier1_price or it.standard_price,
                it.moq or "10 Cartons",
            ])
            ws.row_dimensions[idx].height = 22
            row_fill = zebra_fill if idx % 2 == 0 else PatternFill(fill_type=None)
            for c_idx in range(1, 9):
                c = ws.cell(row=idx, column=c_idx)
                c.border = thin_border
                c.font = Font(name="Arial", size=9)
                if row_fill.fill_type:
                    c.fill = row_fill
                if c_idx in (6, 7):
                    c.number_format = "$#,##0.00"
                    c.alignment = Alignment(horizontal="right", vertical="center")
                elif c_idx in (1, 5, 8):
                    c.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    c.alignment = Alignment(horizontal="left", vertical="center")

        # Auto-fit columns
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    except Exception as exc:
        logger.error("Failed to generate Horeka Excel: %s", exc)
        # Fallback to CSV
        lines = ["Item Code,Category,Product,Packaging,Unit,Standard Price,Bulk Tier 1,MOQ\n"]
        for it in items:
            lines.append(f'"{it.item_code}","{it.category}","{it.product_name}","{it.packaging}","{it.unit}",{it.standard_price},{it.bulk_tier1_price},"{it.moq}"\n')
        return "".join(lines).encode("utf-8")


def attach_horeka_price_list_as_attachment(db: Session, file_format: str = "excel") -> dict[str, Any]:
    """Generates the latest Horeka price list and registers it as an EmailAttachment."""
    date_str = datetime.now().strftime("%Y%m%d")
    if file_format.lower() in ("xlsx", "excel"):
        data = generate_horeka_excel(db)
        filename = f"Kafi_Horeka_B2B_Price_List_{date_str}.xlsx"
        content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        # Fallback / PDF format as Excel for now
        data = generate_horeka_excel(db)
        filename = f"Kafi_Horeka_B2B_Price_List_{date_str}.xlsx"
        content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    registered = register_attachment_from_bytes(data, filename=filename, content_type=content_type)
    return public_attachment(registered)
