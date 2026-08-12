"""Data Synthesis — clean, merge, and dedupe scattered Kafi spreadsheet exports.

Upload multiple XLS/CSV files (all sheets). Rows are mapped to the Master Contacts
layout, company fields are cleaned, duplicates across files are merged, and rows
that already exist in a baseline master file or the CRM are skipped.
"""

from __future__ import annotations

import csv
import io
import re
from collections import defaultdict
from dataclasses import dataclass, field
from typing import Any, Callable, Iterator

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils.exceptions import InvalidFileException

from modules.field_clean import email_dedupe_key, normalize_email_or_empty
from modules.file_to_csv import SUPPORTED_UPLOAD_EXTENSIONS, _resolve_extension, convert_upload_to_csv
from modules.lead_discovery import _dedupe_domain, _normalize_name, parse_csv_candidates

ProgressCallback = Callable[[dict[str, Any]], None]

# ── Company-field cleaning (Usman pass-1 rules, inlined for synthesis) ───────

ZW = re.compile(r"[\u200b\u200c\u200d\ufeff]")
UK_PC = re.compile(r"\b[A-Z]{1,2}\d[A-Z\d]?\s*\d[A-Z]{2}\b", re.I)
CITY_PC = re.compile(
    r"^(?P<body>[A-Za-z][A-Za-z .'\-]{1,40}?)[,\s]+(?P<pc>[A-Z]{1,2}\d[A-Z\d]?\s*\d[A-Z]{2})$",
    re.I,
)
STREET_START = re.compile(
    r"^(?:\d{1,5}[A-Za-z]?\s+|[A-Z]?\d{1,5}\s+).+\b(street|st\.?|road|rd\.?|avenue|ave\.?|"
    r"lane|ln\.?|drive|dr\.?|close|court|ct\.?|place|pl\.?|way|boulevard|blvd|centre|center|"
    r"estate|industrial|park)\b",
    re.I,
)
BARE_LOCATION = re.compile(
    r"^(Middlesex|London|Manchester|Birmingham|Leeds|Glasgow|Edinburgh|Cardiff|Bristol|"
    r"Weston Centre|Industrial Estate|Trading Estate|Business Park)(\b|$).*",
    re.I,
)
EMAILISH = re.compile(r"(?i)^(?:email\s*[·\-:.]?\s*)?.+@.+\..+")
EMAIL_LABEL = re.compile(r"(?i)^email\s*[·\-:]")
ONLY_DASH = re.compile(r"^[\-\—–_./\s]+$")
COMPANYISH = re.compile(
    r"(?i)\b(ltd|limited|llc|inc|corp|company|co\.|plc|pvt|gmbh|sarl|bv|oy|ab|"
    r"trading|foods|distributors?|importers?|exporters?|group|enterprises?|mill|mills)\b"
)


def clean_text(value: object | None) -> str:
    if value is None:
        return ""
    return ZW.sub("", str(value)).replace("\n", " ").strip()


def classify_company(co: str) -> str | None:
    s = co.strip()
    if not s:
        return None
    if ONLY_DASH.match(s) or s in {"---", "N/A", "n/a", "NA", ".", "null", "None"}:
        return "dash_placeholder"
    if EMAIL_LABEL.search(s) or EMAILISH.search(s):
        return "email_as_company"
    if re.fullmatch(r"[A-Z]{1,2}\d[A-Z\d]?\s*\d[A-Z]{2}", s, re.I):
        return "postcode_only"
    if UK_PC.search(s) and not COMPANYISH.search(s):
        return "uk_postcode_as_company"
    if CITY_PC.match(s) and not COMPANYISH.search(s):
        return "city_postcode_as_company"
    if STREET_START.search(s) and not COMPANYISH.search(s):
        return "street_address_as_company"
    if BARE_LOCATION.match(s) and not COMPANYISH.search(s):
        return "location_as_company"
    return None

MASTER_HEADERS = [
    "S. No",
    "Company Name",
    "Business Type",
    "Companies Grading",
    "Designation",
    "Contact Person",
    "Primary Mobile No.",
    "Secondary Mobile No.",
    "Primary Phone No.",
    "Secondary Phone No.",
    "Primary Email",
    "Secondary Email",
    "Country",
    "Product",
    "City",
    "Address",
    "Remarks",
]

MASTER_FIELDS = [
    "serial",
    "company_name",
    "business_type",
    "company_grading",
    "designation",
    "contact_person",
    "primary_mobile",
    "secondary_mobile",
    "primary_phone",
    "secondary_phone",
    "primary_email",
    "secondary_email",
    "country",
    "product",
    "city",
    "address",
    "remarks",
]

EMAIL_RE = re.compile(r"[\w.+-]+@[\w.-]+\.\w+")
PHONE_DIGITS_RE = re.compile(r"\d{8,15}")
PERSON_PREFIX = re.compile(r"^(mr\.?|mrs\.?|ms\.?|miss|dr\.?|prof\.?|shaikh|sheikh)\s+", re.I)

# Column widths for readable Excel export (matches MASTER_HEADERS order).
EXPORT_COLUMN_WIDTHS = (
    8, 28, 18, 14, 14, 22, 16, 16, 16, 16, 28, 28, 14, 14, 14, 32, 24,
)


def _looks_like_person_name(text: str) -> bool:
    """Heuristic: 'Mr. Ali Khan' belongs in Contact Person, not Company Name."""
    s = (text or "").strip()
    if not s or len(s) < 3:
        return False
    if COMPANYISH.search(s) or EMAIL_RE.search(s) or _phone_digits(s):
        return False
    if PERSON_PREFIX.search(s):
        return True
    words = [word for word in re.split(r"\s+", s) if word]
    if not (2 <= len(words) <= 4):
        return False
    if re.search(r"\d", s):
        return False
    business_words = {
        "trading",
        "foods",
        "distributor",
        "distributors",
        "import",
        "export",
        "company",
        "co",
        "ltd",
        "limited",
        "group",
        "center",
        "centre",
        "store",
        "market",
        "hyper",
        "mall",
    }
    lower_words = [word.lower().strip(".,&") for word in words]
    if any(word in business_words for word in lower_words):
        return False
    # Title-case personal names (incl. "Shakeel T.")
    return all(re.match(r"^[A-Z][A-Za-z.'-]*\.?$", word) for word in words)


def _looks_like_company_name(text: str) -> bool:
    s = (text or "").strip()
    if not s or len(s) < 3:
        return False
    if COMPANYISH.search(s):
        return True
    if EMAIL_RE.search(s) or _phone_digits(s):
        return False
    if PERSON_PREFIX.search(s):
        return False
    words = [word for word in re.split(r"\s+", s) if word]
    business_words = {
        "trading",
        "traders",
        "stores",
        "store",
        "agency",
        "agencies",
        "group",
        "llc",
        "ltd",
        "limited",
        "co",
        "holdings",
        "holding",
        "enterprises",
        "enterprise",
        "commercial",
        "international",
        "logistics",
        "distributor",
        "distributors",
        "foods",
        "mills",
        "jewel",
        "jewels",
    }
    lower_words = [word.lower().strip(".,&") for word in words]
    if any(word in business_words for word in lower_words):
        return True
    # e.g. MASAD AL HATMI, GOLDEN FALCON INTERNATIONAL TRADE LLC
    if len(words) >= 2 and s.isupper():
        return True
    return False


def _rebalance_person_company(row: dict[str, str]) -> dict[str, str]:
    company = clean_text(row.get("company_name") or "")
    contact = clean_text(row.get("contact_person") or "")

    if company and contact and _normalize_name(company) == _normalize_name(contact):
        row["contact_person"] = ""
        contact = ""

    if contact and not company and _looks_like_company_name(contact):
        row["company_name"] = contact
        row["contact_person"] = ""
    elif company and not contact and _looks_like_person_name(company):
        row["contact_person"] = company
        row["company_name"] = ""

    return row


def _normalize_header(field: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", field.strip().lower()).strip("_")


def _clean_scalar(value: object | None) -> str:
    if value is None:
        return ""
    text = clean_text(value)
    if re.fullmatch(r"\d+\.0+", text):
        return text.split(".", 1)[0]
    return text


def _phone_digits(value: str) -> str | None:
    digits = re.sub(r"\D", "", value or "")
    if len(digits) < 8 or len(digits) > 15:
        return None
    if len(set(digits)) <= 2:
        return None
    return digits


def _email_key(value: str) -> str | None:
    return email_dedupe_key(value)


def _domain_key(value: str) -> str | None:
    return _dedupe_domain(value.strip() or None)


def _name_country_key(company: str, country: str) -> str | None:
    name = _normalize_name(company or "")
    country_norm = _normalize_name(country or "")
    if not name or len(name) < 3:
        return None
    if country_norm:
        return f"{name}|{country_norm}"
    return name if len(name) >= 5 else None


def row_dedupe_keys(row: dict[str, str]) -> set[str]:
    keys: set[str] = set()
    for phone_field in (
        "primary_mobile",
        "secondary_mobile",
        "primary_phone",
        "secondary_phone",
    ):
        digits = _phone_digits(row.get(phone_field) or "")
        if digits:
            keys.add(f"phone:{digits}")
    for email_field in ("primary_email", "secondary_email"):
        email = _email_key(row.get(email_field) or "")
        if email:
            keys.add(f"email:{email}")
    domain = _domain_key(row.get("website") or row.get("company_name") or "")
    if domain:
        keys.add(f"domain:{domain}")
    nc = _name_country_key(row.get("company_name") or "", row.get("country") or "")
    if nc:
        keys.add(f"name:{nc}")
    return keys


def empty_master_row() -> dict[str, str]:
    return {field: "" for field in MASTER_FIELDS}


def _row_richness(row: dict[str, str]) -> int:
    return sum(1 for value in row.values() if (value or "").strip())


def merge_master_rows(primary: dict[str, str], secondary: dict[str, str]) -> dict[str, str]:
    merged = dict(primary)
    for key, value in secondary.items():
        if key not in merged:
            merged[key] = value
            continue
        if not (merged.get(key) or "").strip() and (value or "").strip():
            merged[key] = value
    return merged


def clean_master_row(row: dict[str, str]) -> dict[str, str]:
    cleaned = dict(row)
    company = clean_text(cleaned.get("company_name") or "")
    reason = classify_company(company) if company else None

    if reason == "email_as_company":
        match = EMAIL_RE.search(company)
        if match and not (cleaned.get("primary_email") or "").strip():
            cleaned["primary_email"] = match.group(0)
        cleaned["company_name"] = ""
    elif reason in {
        "street_address_as_company",
        "uk_postcode_as_company",
        "city_postcode_as_company",
        "location_as_company",
        "postcode_only",
    }:
        if not (cleaned.get("address") or "").strip():
            cleaned["address"] = company
        cleaned["company_name"] = ""
    elif reason == "dash_placeholder":
        cleaned["company_name"] = ""

    company = clean_text(cleaned.get("company_name") or "")
    contact = clean_text(cleaned.get("contact_person") or "")
    if company and not contact and _looks_like_person_name(company):
        cleaned["contact_person"] = company
        cleaned["company_name"] = ""

    for key in MASTER_FIELDS:
        cleaned[key] = _clean_scalar(cleaned.get(key))

    for email_field in ("primary_email", "secondary_email"):
        cleaned[email_field] = normalize_email_or_empty(cleaned.get(email_field))

    return _rebalance_person_company(cleaned)


def _header_map(fieldnames: list[str]) -> dict[str, str | None]:
    """Map master field -> source column name."""

    def col(*names: str) -> str | None:
        wanted = {_normalize_header(name) for name in names}
        for field in fieldnames:
            if _normalize_header(field) in wanted:
                return field
        return None

    has_company_col = bool(
        col(
            "company_name",
            "company",
            "organization",
            "organisation",
            "business_name",
            "firm",
            "client",
            "customer",
            "account",
            "account_name",
        )
    )
    name_col = col("name", "full_name", "person", "contact", "contact_name", "contact_person")
    number_col = col("number", "no", "mobile", "phone", "tel", "contact_number", "whatsapp")
    email_only_col = col(
        "primary_email",
        "email",
        "e_mail",
        "e-mail",
        "mail",
        "contact_email",
    )
    company_col = col(
        "company_name",
        "company",
        "organization",
        "organisation",
        "business_name",
        "business",
        "firm",
        "client",
        "customer",
        "account",
        "account_name",
    )
    if not company_col and name_col and number_col:
        # Phone + name only — name is almost always the contact person.
        company_col = None
    elif not company_col and name_col and email_only_col and not number_col:
        # Name + Email sheets — name is contact person; company stays empty unless mapped.
        company_col = None
    elif not company_col and name_col:
        company_col = name_col
        name_col = None
    elif has_company_col and name_col == company_col:
        name_col = col("contact_person", "contact", "person", "full_name", "representative", "attn")

    phone_from_number = number_col if number_col and not col(
        "primary_mobile_no",
        "primary_mobile",
        "contact_phone",
        "phone",
        "mobile",
    ) else None

    return {
        "serial": col("s_no", "serial_no", "serial", "legacy_serial_no", "no", "sr_no"),
        "company_name": company_col or (fieldnames[0] if fieldnames and not phone_from_number else None),
        "business_type": col(
            "business_type",
            "busniess_type",
            "industry",
            "sector",
            "type",
            "category",
        ),
        "company_grading": col("companies_grading", "company_grading", "grading", "grade"),
        "designation": col("designation", "title", "job_title", "position"),
        "contact_person": name_col
        or col("contact_person", "contact_name", "contact", "person", "full_name", "representative"),
        "primary_mobile": col(
            "primary_mobile_no",
            "primary_mobile",
            "primary_mobile_number",
            "contact_phone",
            "phone",
            "mobile",
            "telephone",
            "tel",
            "cell",
        )
        or phone_from_number,
        "secondary_mobile": col(
            "secondary_mobile_no",
            "secondary_mobile",
            "secondary_mobile_number",
            "alt_mobile",
        ),
        "primary_phone": col(
            "primary_phone_no",
            "primary_phone",
            "primary_telephone",
            "office_phone",
            "landline",
        ),
        "secondary_phone": col(
            "secondary_phone_no",
            "secondary_phone",
            "secondary_telephone",
            "alt_phone",
        ),
        "primary_email": col(
            "primary_email",
            "email",
            "e_mail",
            "e-mail",
            "mail",
            "contact_email",
        ),
        "secondary_email": col(
            "secondary_email",
            "secondary_e_mail",
            "alt_email",
            "alternate_email",
        ),
        "country": col("country", "market", "region", "nation", "location"),
        "product": col(
            "product",
            "products",
            "product_interest",
            "product_type",
            "product_focus",
        ),
        "city": col("city", "town"),
        "address": col("address", "street_address", "full_address"),
        "remarks": col(
            "remarks",
            "remark",
            "notes",
            "note",
            "comment",
            "comments",
            "remarks_02",
            "remarks2",
        ),
        "website": col("website_url", "website", "url", "web", "site", "homepage"),
    }


def _detect_column_role(values: list[str]) -> str | None:
    non_empty = [value.strip() for value in values if value and value.strip()]
    if not non_empty:
        return None
    email_hits = sum(1 for value in non_empty if EMAIL_RE.search(value))
    phone_hits = sum(1 for value in non_empty if _phone_digits(value))
    if email_hits >= max(2, len(non_empty) // 2):
        return "email"
    if phone_hits >= max(2, len(non_empty) // 2):
        return "phone"
    return None


def _apply_column_role_fixes(
    mapping: dict[str, str | None],
    rows: list[dict[str, str]],
) -> dict[str, str | None]:
    fixed = dict(mapping)
    company_col = fixed.get("company_name")
    if not company_col:
        return fixed
    sample_values = [row.get(company_col, "") for row in rows[:40]]
    role = _detect_column_role(sample_values)
    if role == "email" and not fixed.get("primary_email"):
        fixed["primary_email"] = company_col
        fixed["company_name"] = fixed.get("contact_person") or None
    elif role == "phone" and not fixed.get("primary_mobile"):
        fixed["primary_mobile"] = company_col
        fixed["company_name"] = fixed.get("contact_person") or None
    return fixed


def _dict_rows_from_csv(content: str) -> tuple[list[str], list[dict[str, str]]]:
    sample = content[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(content), dialect=dialect)
    fieldnames = [field for field in (reader.fieldnames or []) if field]
    rows: list[dict[str, str]] = []
    for row in reader:
        normalized = {
            (field or "").strip(): _clean_scalar(row.get(field))
            for field in reader.fieldnames or []
            if field
        }
        if any(value.strip() for value in normalized.values()):
            rows.append(normalized)
    return fieldnames, rows


def iter_spreadsheet_tables(
    filename: str | None,
    raw: bytes,
) -> Iterator[tuple[str, str, list[str], list[dict[str, str]]]]:
    """Yield (filename, sheet_name, headers, row dicts) for every sheet/table."""
    display = (filename or "upload").strip() or "upload"
    ext = _resolve_extension(filename, raw)

    if ext == ".csv" or ext == ".tsv":
        content, _ = convert_upload_to_csv(filename, raw)
        headers, rows = _dict_rows_from_csv(content)
        yield display, "Sheet1", headers, rows
        return

    if ext in {".xlsx", ".xlsm"}:
        try:
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except InvalidFileException as exc:
            raise ValueError(
                "Could not read this Excel file. Save it as .xlsx or .csv and upload again."
            ) from exc
        try:
            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                iterator = sheet.iter_rows(values_only=True)
                try:
                    header_row = next(iterator)
                except StopIteration:
                    continue
                headers = [_clean_scalar(cell) for cell in header_row]
                while headers and not headers[-1]:
                    headers.pop()
                if not any(headers):
                    continue
                width = len(headers)
                rows: list[dict[str, str]] = []
                for values in iterator:
                    cells = [_clean_scalar(value) for value in values[:width]]
                    if len(cells) < width:
                        cells.extend([""] * (width - len(cells)))
                    if not any(cell.strip() for cell in cells):
                        continue
                    rows.append(dict(zip(headers, cells)))
                if rows:
                    yield display, sheet_name, headers, rows
        finally:
            workbook.close()
        return

    if ext == ".xls":
        content, _ = convert_upload_to_csv(filename, raw)
        headers, rows = _dict_rows_from_csv(content)
        yield display, "Sheet1", headers, rows
        return

    supported = ", ".join(sorted(SUPPORTED_UPLOAD_EXTENSIONS))
    raise ValueError(f"Unsupported file type. Upload one of: {supported}")


def map_sheet_rows(headers: list[str], rows: list[dict[str, str]]) -> list[dict[str, str]]:
    mapping = _header_map(headers)
    mapping = _apply_column_role_fixes(mapping, rows)
    mapped_rows: list[dict[str, str]] = []

    for source in rows:
        master = empty_master_row()
        for master_field, source_col in mapping.items():
            if master_field == "website" or not source_col:
                continue
            master[master_field] = _clean_scalar(source.get(source_col))
        website = mapping.get("website")
        if website:
            master["website"] = _clean_scalar(source.get(website))
        if not any(master.values()):
            continue
        mapped_rows.append(clean_master_row(master))
    return mapped_rows


def build_existing_key_index(rows: list[dict[str, str]]) -> set[str]:
    keys: set[str] = set()
    for row in rows:
        keys.update(row_dedupe_keys(row))
    return keys


def merge_batch_rows(rows: list[dict[str, str]]) -> tuple[list[dict[str, str]], int]:
    """Union-find merge duplicates within the uploaded batch."""
    if not rows:
        return [], 0

    parent = list(range(len(rows)))

    def find(i: int) -> int:
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i

    def union(a: int, b: int) -> None:
        ra, rb = find(a), find(b)
        if ra != rb:
            parent[rb] = ra

    key_to_index: dict[str, int] = {}
    for index, row in enumerate(rows):
        for key in row_dedupe_keys(row):
            if key in key_to_index:
                union(index, key_to_index[key])
            else:
                key_to_index[key] = index

    clusters: dict[int, list[dict[str, str]]] = defaultdict(list)
    for index, row in enumerate(rows):
        clusters[find(index)].append(row)

    merged: list[dict[str, str]] = []
    merged_count = 0
    for cluster in clusters.values():
        keeper = max(cluster, key=_row_richness)
        for other in cluster:
            if other is not keeper:
                keeper = merge_master_rows(keeper, other)
                merged_count += 1
        merged.append(keeper)
    return merged, merged_count


def load_existing_keys_from_db(db) -> set[str]:
    from db.models import Buyer, Contact

    keys: set[str] = set()
    buyers = db.query(Buyer.id, Buyer.company_name, Buyer.website_url, Buyer.country).all()
    buyer_country = {row.id: (row.country or "") for row in buyers}
    buyer_company = {row.id: (row.company_name or "") for row in buyers}
    for row in buyers:
        pseudo = {
            "company_name": row.company_name or "",
            "country": row.country or "",
            "website": row.website_url or "",
        }
        keys.update(row_dedupe_keys(pseudo))
        domain = _domain_key(row.website_url or "")
        if domain:
            keys.add(f"domain:{domain}")

    contacts = db.query(
        Contact.buyer_id,
        Contact.email,
        Contact.phone,
        Contact.secondary_mobile,
        Contact.primary_phone,
        Contact.secondary_phone,
        Contact.secondary_email,
    ).all()
    for contact in contacts:
        for value in (
            contact.email,
            contact.phone,
            contact.secondary_mobile,
            contact.primary_phone,
            contact.secondary_phone,
            contact.secondary_email,
        ):
            email = _email_key(value or "")
            if email:
                keys.add(f"email:{email}")
            digits = _phone_digits(value or "")
            if digits:
                keys.add(f"phone:{digits}")
        company = buyer_company.get(contact.buyer_id, "")
        country = buyer_country.get(contact.buyer_id, "")
        nc = _name_country_key(company, country)
        if nc:
            keys.add(f"name:{nc}")
    return keys


def export_master_xlsx(rows: list[dict[str, str]]) -> bytes:
    sorted_rows = sorted(
        rows,
        key=lambda row: (
            (row.get("company_name") or row.get("contact_person") or "").strip().lower(),
            (row.get("contact_person") or "").strip().lower(),
            (row.get("primary_email") or "").strip().lower(),
        ),
    )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Master Contacts"
    bold = Font(bold=True)
    sheet.append(MASTER_HEADERS)
    for cell in sheet[1]:
        cell.font = bold
    for col_idx, width in enumerate(EXPORT_COLUMN_WIDTHS, start=1):
        col_letter = sheet.cell(row=1, column=col_idx).column_letter
        sheet.column_dimensions[col_letter].width = width
    sheet.freeze_panes = "A2"

    for index, row in enumerate(sorted_rows, start=1):
        sheet.append(
            [
                str(index),
                row.get("company_name", ""),
                row.get("business_type", ""),
                row.get("company_grading", ""),
                row.get("designation", ""),
                row.get("contact_person", ""),
                row.get("primary_mobile", ""),
                row.get("secondary_mobile", ""),
                row.get("primary_phone", ""),
                row.get("secondary_phone", ""),
                row.get("primary_email", ""),
                row.get("secondary_email", ""),
                row.get("country", ""),
                row.get("product", ""),
                row.get("city", ""),
                row.get("address", ""),
                row.get("remarks", ""),
            ]
        )

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@dataclass
class SynthesisResult:
    output_rows: list[dict[str, str]]
    raw_rows: int = 0
    mapped_rows: int = 0
    merged_duplicates: int = 0
    skipped_existing: int = 0
    sheets_processed: int = 0
    files_processed: int = 0
    messages: list[str] = field(default_factory=list)


def run_synthesis(
    uploads: list[tuple[str | None, bytes]],
    *,
    baseline_uploads: list[tuple[str | None, bytes]] | None = None,
    db=None,
    check_db: bool = True,
    progress: ProgressCallback | None = None,
) -> SynthesisResult:
    result = SynthesisResult(output_rows=[])
    baseline_keys: set[str] = set()
    all_mapped: list[dict[str, str]] = []

    def emit(**fields: Any) -> None:
        if progress:
            progress(fields)

    # Baseline master file(s) — rows here define "already exists"
    baseline_files = baseline_uploads or []
    for filename, raw in baseline_files:
        result.files_processed += 1
        for file_name, sheet_name, headers, rows in iter_spreadsheet_tables(filename, raw):
            result.sheets_processed += 1
            mapped = map_sheet_rows(headers, rows)
            baseline_keys.update(build_existing_key_index(mapped))
            result.messages.append(
                f"Baseline {file_name} / {sheet_name}: indexed {len(mapped)} row(s) for duplicate detection."
            )

    if check_db and db is not None:
        baseline_keys.update(load_existing_keys_from_db(db))
        result.messages.append("Loaded existing CRM contacts for duplicate detection.")

    total_steps = sum(
        1
        for _filename, raw in uploads
        for _file_name, _sheet_name, _headers, rows in iter_spreadsheet_tables(_filename, raw)
        for _row in rows
    )
    processed = 0
    seen_files: set[str] = set()

    for filename, raw in uploads:
        file_label = (filename or "upload").strip() or "upload"
        seen_files.add(file_label)
        for file_name, sheet_name, headers, rows in iter_spreadsheet_tables(filename, raw):
            result.sheets_processed += 1
            result.raw_rows += len(rows)
            mapped = map_sheet_rows(headers, rows)
            result.mapped_rows += len(mapped)

            for row in mapped:
                processed += 1
                keys = row_dedupe_keys(row)
                if keys & baseline_keys:
                    result.skipped_existing += 1
                    emit(
                        processed=processed,
                        total=max(total_steps, 1),
                        phase="deduplicating",
                        current_company=row.get("company_name") or row.get("contact_person") or "",
                    )
                    continue
                all_mapped.append(row)
                emit(
                    processed=processed,
                    total=max(total_steps, 1),
                    phase="cleaning",
                    current_company=row.get("company_name") or row.get("contact_person") or "",
                )

    result.files_processed = len(seen_files) + len(baseline_files)

    emit(processed=total_steps, total=max(total_steps, 1), phase="merging")
    merged_rows, merged_count = merge_batch_rows(all_mapped)
    result.merged_duplicates = merged_count
    result.output_rows = merged_rows
    result.messages.append(
        f"Synthesized {len(merged_rows)} unique row(s) from {result.raw_rows} raw row(s) "
        f"across {result.sheets_processed} sheet(s)."
    )
    if result.skipped_existing:
        result.messages.append(f"Skipped {result.skipped_existing} row(s) already in master/CRM.")
    if merged_count:
        result.messages.append(f"Merged {merged_count} duplicate row(s) within the upload batch.")
    return result


def run_synthesis_via_csv_parser(
    uploads: list[tuple[str | None, bytes]],
    *,
    default_country: str | None = None,
) -> SynthesisResult:
    """Fallback path using the existing CSV candidate parser (single-sheet files)."""
    result = SynthesisResult(output_rows=[])
    rows: list[dict[str, str]] = []
    for filename, raw in uploads:
        content, messages = convert_upload_to_csv(filename, raw)
        result.messages.extend(messages)
        candidates = parse_csv_candidates(content, default_country, include_blank_names=True)
        for candidate in candidates:
            rows.append(
                clean_master_row(
                    {
                        "serial": str(candidate.legacy_serial_no or ""),
                        "company_name": candidate.company_name or "",
                        "business_type": candidate.industry or "",
                        "company_grading": candidate.company_grading or "",
                        "designation": candidate.designation or "",
                        "contact_person": candidate.contact_name or "",
                        "primary_mobile": candidate.phone if candidate.phone != "Not found" else "",
                        "secondary_mobile": candidate.secondary_mobile or "",
                        "primary_phone": candidate.primary_phone or "",
                        "secondary_phone": candidate.secondary_phone or "",
                        "primary_email": candidate.email if candidate.email != "Not found" else "",
                        "secondary_email": candidate.secondary_email or "",
                        "country": candidate.country or "",
                        "product": candidate.product_interest or "",
                        "city": candidate.city or "",
                        "address": candidate.address or "",
                        "remarks": candidate.remarks or "",
                        "website": candidate.website_url or "",
                    }
                )
            )
    merged_rows, merged_count = merge_batch_rows(rows)
    result.output_rows = merged_rows
    result.raw_rows = len(rows)
    result.mapped_rows = len(rows)
    result.merged_duplicates = merged_count
    return result
