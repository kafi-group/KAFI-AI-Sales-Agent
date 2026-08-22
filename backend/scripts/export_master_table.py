"""Export full Master Table (all buyers) from production DB before wipe/re-import.

Usage (from backend/):
  python scripts/export_master_table.py

Outputs under:
  D:\\First AI World\\Cursor Projects\\sales, social, hr Aug 2026\\master-data-backup\\
"""
from __future__ import annotations

import os
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

_BACKEND_DIR = Path(__file__).resolve().parent.parent
if str(_BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(_BACKEND_DIR))
os.chdir(_BACKEND_DIR)

from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import create_engine, text
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import NullPool

from modules.data_synthesis import EXPORT_COLUMN_WIDTHS, MASTER_HEADERS

OUT_DIR = Path(r"D:\First AI World\Cursor Projects\sales, social, hr Aug 2026\master-data-backup")
ZW = re.compile(r"[\u200b\u200c\u200d\ufeff]")

EXTENDED_HEADERS = [
    "buyer_id",
    "legacy_serial_no",
    *MASTER_HEADERS,
    "Website",
    "Source",
    "Assigned To",
    "Created At",
]


def clean_text(value: object) -> str:
    if value is None:
        return ""
    return ZW.sub("", str(value)).replace("\n", " ").strip()


def fetch_rows(db: Session) -> list[dict]:
    sql = text(
        """
        SELECT
          b.id AS buyer_id,
          b.legacy_serial_no,
          b.company_name,
          b.industry AS business_type,
          b.company_grading,
          c.designation AS designation,
          c.full_name AS contact_person,
          c.phone AS primary_mobile,
          c.secondary_mobile AS secondary_mobile,
          c.primary_phone AS primary_phone,
          c.secondary_phone AS secondary_phone,
          c.email AS primary_email,
          c.secondary_email AS secondary_email,
          b.country,
          b.product_interest AS product,
          b.city,
          b.address,
          b.remarks,
          b.website_url AS website,
          b.source,
          b.assigned_to,
          b.created_at
        FROM buyers b
        LEFT JOIN LATERAL (
          SELECT *
          FROM contacts c
          WHERE c.buyer_id = b.id
          ORDER BY c.id ASC
          LIMIT 1
        ) c ON TRUE
        ORDER BY COALESCE(b.legacy_serial_no, b.id) ASC
        """
    )
    result = db.execute(sql)
    cols = list(result.keys())
    return [dict(zip(cols, row)) for row in result.fetchall()]


def row_to_reimport(r: dict, index: int) -> list:
    return [
        str(index),
        clean_text(r.get("company_name")),
        clean_text(r.get("business_type")),
        clean_text(r.get("company_grading")),
        clean_text(r.get("designation")),
        clean_text(r.get("contact_person")),
        clean_text(r.get("primary_mobile")),
        clean_text(r.get("secondary_mobile")),
        clean_text(r.get("primary_phone")),
        clean_text(r.get("secondary_phone")),
        clean_text(r.get("primary_email")),
        clean_text(r.get("secondary_email")),
        clean_text(r.get("country")),
        clean_text(r.get("product")),
        clean_text(r.get("city")),
        clean_text(r.get("address")),
        clean_text(r.get("remarks")),
    ]


def row_to_extended(r: dict, index: int) -> list:
    created = r.get("created_at")
    created_str = created.isoformat() if created is not None else ""
    return [
        r.get("buyer_id"),
        r.get("legacy_serial_no"),
        *row_to_reimport(r, index),
        clean_text(r.get("website")),
        clean_text(r.get("source")),
        clean_text(r.get("assigned_to")),
        created_str,
    ]


def write_sheet(ws, headers: list[str], rows: list[list], *, column_widths: tuple[int, ...] | None = None) -> None:
    bold = Font(bold=True)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = bold
    if column_widths:
        for col_idx, width in enumerate(column_widths, start=1):
            col_letter = ws.cell(row=1, column=col_idx).column_letter
            ws.column_dimensions[col_letter].width = width
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"


def make_engine(database_url: str):
    url = database_url
    if "pooler.supabase.com:5432" in url:
        url = url.replace(":5432/", ":6543/")
    return create_engine(url, poolclass=NullPool, pool_pre_ping=True)


def main() -> None:
    database_url = (os.environ.get("DATABASE_URL") or "").strip()
    if not database_url:
        load_dotenv(_BACKEND_DIR / ".env")
        database_url = (os.environ.get("DATABASE_URL") or "").strip()
    if not database_url:
        raise SystemExit("DATABASE_URL is not set. Add it to backend/.env or run via railway run.")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S")
    reimport_path = OUT_DIR / f"Master-Contacts-REIMPORT-{stamp}.xlsx"
    backup_path = OUT_DIR / f"Master-Table-FULL-BACKUP-{stamp}.xlsx"

    engine = make_engine(database_url)
    SessionLocal = sessionmaker(bind=engine)
    with SessionLocal() as db:
        raw = fetch_rows(db)

    reimport_rows = [row_to_reimport(r, i) for i, r in enumerate(raw, start=1)]
    extended_rows = [row_to_extended(r, i) for i, r in enumerate(raw, start=1)]

    wb_reimport = Workbook()
    ws = wb_reimport.active
    ws.title = "Master Contacts"
    write_sheet(ws, MASTER_HEADERS, reimport_rows, column_widths=EXPORT_COLUMN_WIDTHS)
    wb_reimport.save(reimport_path)

    wb_backup = Workbook()
    ws_full = wb_backup.active
    ws_full.title = "Master Table Full"
    write_sheet(ws_full, EXTENDED_HEADERS, extended_rows)
    wb_backup.save(backup_path)

    by_source: dict[str, int] = {}
    for r in raw:
        key = clean_text(r.get("source")) or "(empty)"
        by_source[key] = by_source.get(key, 0) + 1

    print(f"Exported {len(raw)} master rows")
    print(f"Re-import file: {reimport_path}")
    print(f"Full backup:    {backup_path}")
    print("Rows by source:")
    for source, count in sorted(by_source.items(), key=lambda x: (-x[1], x[0])):
        print(f"  {source}: {count}")


if __name__ == "__main__":
    main()
