"""Pass-1 cleaner for Master Contacts file.

Moves address-like Company Name values into Address when appropriate.
Produces:
  - cleaned workbook
  - review sheet of changes
  - updates process log
"""
from __future__ import annotations

import re
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

SRC = Path(r"c:\Users\Abc\Downloads\Master Contacts file.xlsx")
OUT_DIR = Path(r"D:\First AI World\Cursor Projects\sales, social, hr Aug 2026")
OUT_XLSX = OUT_DIR / "Master-Contacts-CLEANED-pass1.xlsx"
LOG = OUT_DIR / "Master-Contacts-cleaning-log.md"

UK_PC = re.compile(r"\b[A-Z]{1,2}\d[A-Z\d]?\s*\d[A-Z]{2}\b", re.I)
CITY_PC = re.compile(
    r"^(?P<body>[A-Za-z][A-Za-z .'\-]{1,40}?)[,\s]+(?P<pc>[A-Z]{1,2}\d[A-Z\d]?\s*\d[A-Z]{2})$",
    re.I,
)
STREET_START = re.compile(
    r"^(?:\d{1,5}[A-Za-z]?\s+|[A-Z]?\d{1,5}\s+).+\b(street|st\.?|road|rd\.?|avenue|ave\.?|"
    r"lane|ln\.?|drive|dr\.?|close|court|ct\.?|place|pl\.?|way|boulevard|blvd|centre|center|"
    r"estate|industrial|park)\b",
    re.I,
)
BARE_LOCATION = re.compile(
    r"^(Middlesex|London|Manchester|Birmingham|Leeds|Glasgow|Edinburgh|Cardiff|Bristol|"
    r"Weston Centre|Industrial Estate|Trading Estate|Business Park)(\b|$).*",
    re.I,
)
EMAILISH = re.compile(r"(?i)^(?:email\s*[·\-:.]?\s*)?.+@.+\..+")
EMAIL_LABEL = re.compile(r"(?i)^email\s*[·\-:]")
ONLY_DASH = re.compile(r"^[\-\—–_./\s]+$")
COMPANYISH = re.compile(
    r"(?i)\b(ltd|limited|llc|inc|corp|company|co\.|plc|pvt|gmbh|sarl|bv|oy|ab|"
    r"trading|foods|distributors?|importers?|exporters?|group|enterprises?|mill|mills)\b"
)
ZW = re.compile(r"[\u200b\u200c\u200d\ufeff]")


def clean_text(value: object) -> str:
    if value is None:
        return ""
    return ZW.sub("", str(value)).replace("\n", " ").strip()


def classify_company(co: str) -> str | None:
    s = co.strip()
    if not s:
        return None
    if ONLY_DASH.match(s) or s in {"---", "N/A", "n/a", "NA", ".", "null", "None"}:
        return "dash_placeholder"
    if EMAIL_LABEL.search(s) or EMAILISH.search(s):
        return "email_as_company"
    if re.fullmatch(r"[A-Z]{1,2}\d[A-Z\d]?\s*\d[A-Z]{2}", s, re.I):
        return "postcode_only"
    if UK_PC.search(s) and not COMPANYISH.search(s):
        return "uk_postcode_as_company"
    if CITY_PC.match(s) and not COMPANYISH.search(s):
        return "city_postcode_as_company"
    if STREET_START.search(s) and not COMPANYISH.search(s):
        return "street_address_as_company"
    if BARE_LOCATION.match(s) and not COMPANYISH.search(s):
        return "location_as_company"
    return None


def main() -> None:
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb.active
    ws.title = "Cleaned"

    # Normalize header row
    headers: list[str] = []
    for col in range(1, ws.max_column + 1):
        raw = clean_text(ws.cell(1, col).value)
        if raw.lower() == "busniess type":
            raw = "Business Type"
        headers.append(raw)
        ws.cell(1, col).value = raw

    col_index = {h: i + 1 for i, h in enumerate(headers)}
    c_company = col_index["Company Name"]
    c_address = col_index["Address"]
    c_city = col_index.get("City")
    c_email = col_index.get("Primary Email")
    c_country = col_index.get("Country")

    review = wb.create_sheet("Needs Review")
    review.append(
        [
            "Excel Row",
            "Rule",
            "Old Company Name",
            "New Company Name",
            "Old Address",
            "New Address",
            "City",
            "Country",
            "Notes",
        ]
    )
    for cell in review[1]:
        cell.font = Font(bold=True)

    changes = 0
    for row in range(2, ws.max_row + 1):
        company = clean_text(ws.cell(row, c_company).value)
        address = clean_text(ws.cell(row, c_address).value)
        city = clean_text(ws.cell(row, c_city).value) if c_city else ""
        country = clean_text(ws.cell(row, c_country).value) if c_country else ""
        email = clean_text(ws.cell(row, c_email).value) if c_email else ""

        # Always strip ZW chars from company/address
        if ws.cell(row, c_company).value is not None:
            ws.cell(row, c_company).value = company or None
        if ws.cell(row, c_address).value is not None:
            ws.cell(row, c_address).value = address or None

        reason = classify_company(company)
        if not reason:
            continue

        old_company, old_address = company, address
        new_company, new_address = company, address
        notes = ""

        if reason == "dash_placeholder":
            new_company = ""
            notes = "Cleared placeholder company"
        elif reason == "email_as_company":
            # Extract email if present
            m = re.search(r"[\w.+-]+@[\w.-]+\.\w+", company)
            extracted = m.group(0) if m else ""
            if extracted and c_email and not email:
                ws.cell(row, c_email).value = extracted
                notes = f"Moved email to Primary Email ({extracted})"
            else:
                notes = "Cleared email-like company; Primary Email already set or no extractable email"
            new_company = ""
            if not address:
                # keep nothing in address from email labels
                pass
        else:
            # Address-like company
            if not address:
                new_address = company
                notes = "Moved address-like Company Name → Address"
            else:
                notes = "Address already filled; cleared address-like Company Name"
            new_company = ""

        ws.cell(row, c_company).value = new_company or None
        ws.cell(row, c_address).value = new_address or None
        review.append(
            [
                row,
                reason,
                old_company,
                new_company,
                old_address,
                new_address,
                city,
                country,
                notes,
            ]
        )
        changes += 1

    wb.save(OUT_XLSX)
    print(f"Saved {OUT_XLSX}")
    print(f"Changes logged on Needs Review: {changes}")

    stamp = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    appendix = [
        "",
        f"## Pass 1 applied ({stamp})",
        "",
        f"- Output: `{OUT_XLSX.name}`",
        f"- Rows changed: **{changes}**",
        "- Sheet `Needs Review` lists every automatic change for Usman.",
        "- Company names that were addresses were moved to Address (or cleared if Address already had a value).",
        "- No company names were invented.",
        "",
    ]
    if LOG.exists():
        LOG.write_text(LOG.read_text(encoding="utf-8") + "\n".join(appendix), encoding="utf-8")
    else:
        LOG.write_text("# Master Contacts cleaning log\n" + "\n".join(appendix), encoding="utf-8")
    print(f"Updated {LOG}")


if __name__ == "__main__":
    main()
