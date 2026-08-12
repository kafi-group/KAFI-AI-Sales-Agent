"""Clean a Master Contacts XLSX for safe Old clients upload.

Fixes:
- Strip Excel/text apostrophes from email cells (e.g. info@arcadie.fr')
- Move phone-only values out of email columns into phone fields
- Extract embedded emails from messy cells
- Company/contact rebalance + address fixes (data_synthesis rules)
- Dedupe within file (phone → email → name+country)
- Sort by Country, Company Name; renumber S. No

Usage:
  python scripts/clean_master_contacts_xlsx.py "path/to/file.xlsx"
"""
from __future__ import annotations

import re
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

# Run from backend/
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from modules.data_synthesis import (  # noqa: E402
    EMAIL_RE,
    EXPORT_COLUMN_WIDTHS,
    MASTER_FIELDS,
    MASTER_HEADERS,
    _clean_scalar,
    _email_key,
    _phone_digits,
    clean_master_row,
    merge_master_rows,
    row_dedupe_keys,
)

EMAIL_WRAP_QUOTES = re.compile(r"^['\"`]+|['\"`]+$")
PHONEISH = re.compile(r"^[\d\s+\-().]+$")


def clean_email_cell(value: str) -> tuple[str, str | None]:
    """Return (cleaned_email, note). Moves phone-only values out (empty email)."""
    raw = _clean_scalar(value)
    if not raw:
        return "", None
    stripped = EMAIL_WRAP_QUOTES.sub("", raw.strip()).strip()
    stripped = stripped.strip("'\"`")
    if not stripped:
        return "", "cleared_empty_email"
    if "@" not in stripped and PHONEISH.match(stripped) and _phone_digits(stripped):
        return "", "phone_in_email_column"
    match = EMAIL_RE.search(stripped)
    if match:
        email = match.group(0).lower()
        if email != raw.strip().lower():
            return email, "fixed_email_format"
        return email, None
    if "@" in stripped:
        # Has @ but failed regex — strip quotes and keep best effort
        fixed = EMAIL_WRAP_QUOTES.sub("", stripped).strip("'\"`").lower()
        if EMAIL_RE.search(fixed):
            return EMAIL_RE.search(fixed).group(0).lower(), "fixed_email_format"
    return "", "invalid_email_cleared"


def row_from_values(headers: list[str], values: tuple) -> dict[str, str]:
    header_index = {h: i for i, h in enumerate(headers)}
    field_to_header = {
        "serial": "S. No",
        "company_name": "Company Name",
        "business_type": "Business Type",
        "company_grading": "Companies Grading",
        "designation": "Designation",
        "contact_person": "Contact Person",
        "primary_mobile": "Primary Mobile No.",
        "secondary_mobile": "Secondary Mobile No.",
        "primary_phone": "Primary Phone No.",
        "secondary_phone": "Secondary Phone No.",
        "primary_email": "Primary Email",
        "secondary_email": "Secondary Email",
        "country": "Country",
        "product": "Product",
        "city": "City",
        "address": "Address",
        "remarks": "Remarks",
    }
    row: dict[str, str] = {}
    for field, header in field_to_header.items():
        idx = header_index.get(header)
        if idx is None or idx >= len(values):
            row[field] = ""
        else:
            row[field] = _clean_scalar(values[idx])
    return row


def row_to_values(row: dict[str, str]) -> list[str]:
    return [row.get(field, "") or "" for field in MASTER_FIELDS]


def apply_email_fixes(row: dict[str, str]) -> list[str]:
    notes: list[str] = []
    for email_field, phone_field in (
        ("primary_email", "primary_mobile"),
        ("secondary_email", "secondary_phone"),
    ):
        cleaned, note = clean_email_cell(row.get(email_field) or "")
        if note == "phone_in_email_column":
            digits = _phone_digits(row.get(email_field) or "")
            if digits and not (row.get(phone_field) or "").strip():
                row[phone_field] = row.get(email_field) or ""
            row[email_field] = ""
            notes.append(f"{email_field}: moved phone to {phone_field}")
        elif note:
            row[email_field] = cleaned
            if note != "cleared_empty_email":
                notes.append(f"{email_field}: {note}")
        else:
            row[email_field] = cleaned
    return notes


def fix_phone_as_company(row: dict[str, str]) -> str | None:
    company = (row.get("company_name") or "").strip()
    if not company:
        return None
    if EMAIL_RE.search(company):
        return None
    digits = _phone_digits(company)
    if not digits or not re.fullmatch(r"[\d\s+\-().]+", company):
        return None
    if not (row.get("primary_mobile") or "").strip():
        row["primary_mobile"] = company
    row["company_name"] = ""
    return "moved_phone_from_company_name"


def finalize_row_for_import(row: dict[str, str]) -> tuple[dict[str, str], str | None]:
    """Ensure every import row has a real company name (never blank → Unnamed)."""
    company = (row.get("company_name") or "").strip()
    contact = (row.get("contact_person") or "").strip()
    mobile = (
        (row.get("primary_mobile") or "").strip()
        or (row.get("secondary_mobile") or "").strip()
        or (row.get("primary_phone") or "").strip()
        or (row.get("secondary_phone") or "").strip()
    )
    email = (row.get("primary_email") or "").strip()

    if company:
        return row, None

    if contact:
        row["company_name"] = contact
        return row, "used_contact_as_company"

    if email:
        match = EMAIL_RE.search(email)
        if match:
            row["company_name"] = match.group(0)
            return row, "used_email_as_company"

    if mobile:
        row["company_name"] = mobile
        return row, "used_phone_as_company"

    return row, None


def is_importable_row(row: dict[str, str]) -> bool:
    """Drop country-only / grading-only junk that would import as Unnamed."""
    company = (row.get("company_name") or "").strip()
    contact = (row.get("contact_person") or "").strip()
    mobile = (
        (row.get("primary_mobile") or "").strip()
        or (row.get("secondary_mobile") or "").strip()
        or (row.get("primary_phone") or "").strip()
        or (row.get("secondary_phone") or "").strip()
    )
    email = (row.get("primary_email") or "").strip()
    if company or contact or mobile or email:
        return True
    return False


def is_meaningful_row(row: dict[str, str]) -> bool:
    keys = (
        "company_name",
        "contact_person",
        "primary_mobile",
        "primary_phone",
        "primary_email",
        "country",
        "city",
        "address",
    )
    return any((row.get(k) or "").strip() for k in keys)


def dedupe_rows(rows: list[dict[str, str]]) -> tuple[list[dict[str, str]], int]:
    index: dict[str, int] = {}
    merged: list[dict[str, str]] = []
    dupes = 0
    for row in rows:
        keys = row_dedupe_keys(row)
        hit = next((index[k] for k in keys if k in index), None)
        if hit is not None:
            merged[hit] = merge_master_rows(merged[hit], row)
            dupes += 1
        else:
            pos = len(merged)
            merged.append(row)
            for k in keys:
                index[k] = pos
    return merged, dupes


def sort_rows(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    return sorted(
        rows,
        key=lambda r: (
            (r.get("country") or "").lower(),
            (r.get("company_name") or "").lower(),
            (r.get("contact_person") or "").lower(),
            (r.get("primary_email") or "").lower(),
        ),
    )


def write_workbook(path: Path, rows: list[dict[str, str]], review: list[list]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Master Contacts"
    ws.append(MASTER_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, row in enumerate(rows, start=1):
        row["serial"] = str(i)
        ws.append(row_to_values(row))
    ws.freeze_panes = "A2"
    for col_idx, width in enumerate(EXPORT_COLUMN_WIDTHS, start=1):
        letter = ws.cell(row=1, column=col_idx).column_letter
        ws.column_dimensions[letter].width = width
    if review:
        rev = wb.create_sheet("Cleaning log")
        rev.append(["Row", "Field", "Before", "After", "Note"])
        for cell in rev[1]:
            cell.font = Font(bold=True)
        for line in review:
            rev.append(line)
    wb.save(path)


def process(path: Path) -> dict[str, object]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = [str(h).strip() if h is not None else "" for h in next(ws.iter_rows(max_row=1, values_only=True))]
    raw_rows: list[dict[str, str]] = []
    review: list[list] = []
    stats = defaultdict(int)

    for rno, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        row = row_from_values(headers, values)
        if not is_meaningful_row(row):
            stats["empty_skipped"] += 1
            continue
        before_emails = {
            "primary_email": row.get("primary_email") or "",
            "secondary_email": row.get("secondary_email") or "",
        }
        email_notes = apply_email_fixes(row)
        for field in ("primary_email", "secondary_email"):
            before = before_emails[field]
            after = row.get(field) or ""
            if before != after:
                stats["email_cells_fixed"] += 1
                review.append([rno, field, before, after, "; ".join(email_notes) or "cleaned"])
        cleaned = clean_master_row(row)
        note = fix_phone_as_company(cleaned)
        if note:
            review.append([rno, "company_name", row.get("company_name") or "", cleaned.get("company_name") or "", note])
            stats["company_phone_fixed"] += 1
        cleaned, promote_note = finalize_row_for_import(cleaned)
        if promote_note:
            review.append([rno, "company_name", "", cleaned.get("company_name") or "", promote_note])
            stats["company_promoted"] += 1
        if not is_importable_row(cleaned):
            stats["junk_skipped"] += 1
            review.append([rno, "ROW", "kept", "skipped", "country_or_grading_only"])
            continue
        raw_rows.append(cleaned)
        stats["rows_read"] += 1

    wb.close()

    deduped, dupes = dedupe_rows(raw_rows)
    stats["dupes_merged"] = dupes
    sorted_rows = sort_rows(deduped)
    stats["rows_output"] = len(sorted_rows)

    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    out_dir = path.parent
    out_path = out_dir / f"{path.stem} - CLEANED {stamp}.xlsx"
    write_workbook(out_path, sorted_rows, review)

    return {
        "input": str(path),
        "output": str(out_path),
        "stats": dict(stats),
        "review_samples": review[:15],
    }


def main() -> None:
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    path = Path(sys.argv[1]).expanduser().resolve()
    if not path.is_file():
        raise SystemExit(f"File not found: {path}")
    result = process(path)
    print("CLEANING COMPLETE")
    print(f"Input:  {result['input']}")
    print(f"Output: {result['output']}")
    for k, v in result["stats"].items():
        print(f"  {k}: {v}")
    if result["review_samples"]:
        print("\nSample fixes:")
        for line in result["review_samples"]:
            print(f"  row {line[0]} {line[1]}: {line[2]!r} -> {line[3]!r} ({line[4]})")


if __name__ == "__main__":
    main()
